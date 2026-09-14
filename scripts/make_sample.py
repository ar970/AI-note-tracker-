"""Build a sample deck + transcript so the pipeline can be exercised without API keys.

    python scripts/make_sample.py sample/
    python -m classnotes run --deck sample/deck.pptx \\
        --transcript sample/transcript.json --provider fake --session-id sample

The transcript is deliberately written the way real classroom ASR comes out:
Hinglish code-switching, and technical terms mangled the way a distant mic
mangles them ("market segment nation", "psycho graphic", "the wells framework"
for VALS). That is what step [3] exists to repair, so a clean sample would
prove nothing.
"""
import json, sys
from pathlib import Path
from pptx import Presentation

out = Path(sys.argv[1])
out.mkdir(parents=True, exist_ok=True)

SLIDES = [
    ("Market Segmentation", ["Dividing a heterogeneous market", "Four bases of segmentation", "Geographic and Demographic"]),
    ("Psychographic Segmentation", ["Lifestyle and values", "VALS Framework", "Harder to measure than Demographic"]),
    ("Targeting Strategies", ["Undifferentiated Marketing", "Differentiated Marketing", "Concentrated Marketing"]),
    ("Positioning", ["Perceptual Mapping", "Points of Parity and Points of Difference", "Positioning Statement"]),
    ("Price Elasticity", ["Elastic versus Inelastic Demand", "Cross Price Elasticity"]),
]

prs = Presentation()
layout = prs.slide_layouts[1]
for title, bullets in SLIDES:
    slide = prs.slides.add_slide(layout)
    slide.shapes.title.text = title
    body = slide.placeholders[1].text_frame
    body.text = bullets[0]
    for b in bullets[1:]:
        body.add_paragraph().text = b
    slide.notes_slide.notes_text_frame.text = f"Spend time on {title.lower()}."
prs.save(out / "deck.pptx")

# Deliberately mangled, Hinglish, like real classroom ASR output.
LINES = [
    (0,   "Okay so aaj we will start with market segmentation, theek hai?"),
    (25,  "Market segment nation basically means dividing a heterogenous market."),
    (70,  "There are four basis, geographic, demographic, psycho graphic and behavioural."),
    (140, "Psycho graphic is about lifestyle and values, the wells framework."),
    (210, "Ye thoda difficult hai to measure compared to demographic."),
    (260, "Now targeting. Three strategies - undifferentiated, differentiated, concentrated."),
    (330, "Positioning, perceptual mapping, points of parity aur points of difference."),
    (400, "Price elasticity - elastic versus inelastic demand. Ye exam mein aayega, note kar lo."),
    (430, "Cross price elasticity, chalo quickly, we are out of time."),
]
segments = [{"start": float(s), "end": float(s + 20), "text": t} for s, t in LINES]
(out / "transcript.json").write_text(json.dumps({
    "language": "english", "duration_seconds": 460.0,
    "model": "whisper-large-v3-turbo", "segments": segments,
}, indent=2))
print(f"wrote {out}/deck.pptx and {out}/transcript.json")
