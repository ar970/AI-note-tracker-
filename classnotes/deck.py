"""[2] Slide extraction.

The deck is worth more than it looks. It gives you:

  - the correct spelling of every technical term (ASR will not get these right)
  - the lecture's intended structure, free — slide titles are section headers
  - an ordering to anchor the transcript against

So this module runs *before* transcription as well as after: its `terms()`
output seeds Whisper's decoder prompt.
"""

from __future__ import annotations

import re
from dataclasses import dataclass
from pathlib import Path

ACRONYM = re.compile(r"\b[A-Z]{2,6}\b")
# Capitalised phrases, allowing the lowercase connectors that hold real terms
# together ("Points of Parity", "Return on Investment"). Capped at four words.
TITLE_PHRASE = re.compile(
    r"\b[A-Z][a-z]{2,}(?:\s+(?:of|on|in|and|for|the|to|per|vs)?\s*[A-Z][a-z]{2,}){0,3}\b"
)
STOPWORDS = {
    "The", "And", "For", "With", "That", "This", "From", "Into", "Your",
    "What", "When", "How", "Why", "Who", "All", "Any", "Our", "Their",
    "Slide", "Lecture", "Chapter", "Unit", "Page", "Introduction", "Agenda",
    "Note", "Example", "Summary", "Overview", "Today", "Next", "Recap",
}

# The acronym pattern matches any 2-6 letter all-caps run, which on a deck with
# ALL-CAPS HEADINGS — the norm in a lot of Indian university decks — means THE,
# OF and AND get matched *and* given the acronym weight bonus, outranking the
# real jargon in a budget that only fits about twenty terms. Ordinary words
# caught this way (MARKET, VALUE) are harmless and often genuinely the deck's
# vocabulary; pure function words and date fragments are not.
UPPERCASE_NOISE = {
    "THE", "AND", "FOR", "WITH", "THAT", "THIS", "FROM", "INTO", "YOUR", "ARE",
    "WAS", "WERE", "NOT", "BUT", "ITS", "WHAT", "WHEN", "HOW", "WHY", "WHO",
    "ALL", "ANY", "OUR", "THEN", "THAN", "THEM", "THESE", "THOSE", "BEEN",
    "ONLY", "OVER", "SUCH", "MORE", "MOST", "ALSO", "SEE", "USE", "NEW", "END",
    "ONE", "TWO", "OF", "TO", "IN", "ON", "AT", "BY", "OR", "AS", "IS", "IT",
    "BE", "AN", "IF", "SO", "NO", "DO", "WE", "YOU", "MAY", "CAN", "WILL",
    "HAS", "HAVE", "UNIT", "PART", "PAGE", "NOTE",
    "JAN", "FEB", "MAR", "APR", "JUN", "JUL", "AUG", "SEP", "OCT", "NOV", "DEC",
    "JUNE", "JULY", "MARCH", "APRIL",
}


@dataclass
class Slide:
    number: int
    title: str
    body: list[str]
    notes: str = ""

    def as_markdown(self) -> str:
        head = f"## Slide {self.number}" + (f" — {self.title}" if self.title else "")
        lines = [head]
        lines += [f"- {line}" for line in self.body]
        if self.notes:
            lines.append(f"> Speaker notes: {self.notes}")
        return "\n".join(lines)


@dataclass
class Deck:
    slides: list[Slide]
    source: str

    def as_markdown(self) -> str:
        return "\n\n".join(slide.as_markdown() for slide in self.slides)

    @property
    def text_characters(self) -> int:
        return sum(
            len(s.title) + sum(len(line) for line in s.body) for s in self.slides
        )

    @property
    def looks_scanned(self) -> bool:
        """True when there is far too little text for the number of slides.

        A deck exported as images — scanned handouts, or a PDF of photographed
        slides — parses without error and yields almost nothing. Left unsaid,
        that silently degrades alignment to transcript-only and the notes come
        out vague for a reason nobody can see.
        """
        if not self.slides:
            return True
        return self.text_characters / len(self.slides) < 20

    def terms(self, limit: int = 60) -> list[str]:
        """Candidate technical vocabulary, ranked by how deck-specific it looks.

        Whisper's decoder prompt is ~224 tokens, so this list is a small budget
        that has to be spent well. Two rules keep it clean:

        - **Single words only count from slide titles.** Every bullet begins
          with a capital letter, so matching single capitalised words anywhere
          fills the budget with "Four", "Dividing", "Harder" — words ASR was
          never going to get wrong, crowding out the terms it will.
        - **Multi-word phrases and acronyms count anywhere.** These are exactly
          what a distant mic mangles, and they are unambiguously deck-specific.
        """
        scores: dict[str, float] = {}

        def add(text: str, weight: float, allow_single: bool) -> None:
            for match in ACRONYM.findall(text):
                if match in UPPERCASE_NOISE:
                    continue
                scores[match] = scores.get(match, 0.0) + weight * 1.5
            for match in TITLE_PHRASE.findall(text):
                match = " ".join(match.split())
                words = match.split()
                if len(match) < 4 or words[0] in STOPWORDS:
                    continue
                if len(words) == 1 and not allow_single:
                    continue
                scores[match] = scores.get(match, 0.0) + weight

        for slide in self.slides:
            add(slide.title, 3.0, allow_single=True)
            for line in slide.body:
                add(line, 1.0, allow_single=False)
            if slide.notes:
                add(slide.notes, 0.5, allow_single=False)

        ranked = sorted(scores, key=lambda t: (-scores[t], t))
        return ranked[:limit]


def _clean(lines: list[str]) -> list[str]:
    out: list[str] = []
    for raw in lines:
        line = re.sub(r"\s+", " ", raw).strip(" \t•-–—·")
        if len(line) > 1:
            out.append(line)
    return out


def _from_pdf(path: Path) -> Deck:
    from pypdf import PdfReader
    from pypdf.errors import DependencyError

    reader = PdfReader(str(path))

    if reader.is_encrypted:
        # Most "encrypted" lecture decks are not password-protected at all:
        # faculty export with copy/print restrictions, which encrypts the file
        # with an *empty* user password. pypdf still refuses to read it until
        # you say so. Without this, a student uploads the professor's PDF and
        # the pilot dies on a FileNotDecryptedError nobody can interpret.
        try:
            unlocked = reader.decrypt("")
        except DependencyError as exc:
            raise ValueError(
                f"{path.name} uses AES encryption and the 'cryptography' package "
                f"is missing. Run: pip install -r requirements.txt ({exc})"
            ) from exc
        if not unlocked:
            raise ValueError(
                f"{path.name} is password-protected. Ask whoever shared it for an "
                "unlocked copy, or export the slides to PDF again without a password."
            )

    slides: list[Slide] = []
    for index, page in enumerate(reader.pages, start=1):
        lines = _clean((page.extract_text() or "").splitlines())
        title = lines[0] if lines else ""
        slides.append(Slide(number=index, title=title, body=lines[1:]))
    return Deck(slides=slides, source=path.name)


def _from_pptx(path: Path) -> Deck:
    from pptx import Presentation

    presentation = Presentation(str(path))
    slides: list[Slide] = []
    for index, slide in enumerate(presentation.slides, start=1):
        title = ""
        if slide.shapes.title is not None and slide.shapes.title.has_text_frame:
            title = slide.shapes.title.text.strip()

        body: list[str] = []
        for shape in slide.shapes:
            if shape is slide.shapes.title or not shape.has_text_frame:
                continue
            body.extend(p.text for p in shape.text_frame.paragraphs)

        notes = ""
        if slide.has_notes_slide and slide.notes_slide.notes_text_frame is not None:
            notes = slide.notes_slide.notes_text_frame.text.strip()

        slides.append(
            Slide(number=index, title=title, body=_clean(body), notes=" ".join(notes.split()))
        )
    return Deck(slides=slides, source=path.name)


def load(path: Path) -> Deck:
    suffix = path.suffix.lower()
    if suffix == ".pdf":
        return _from_pdf(path)
    if suffix in {".pptx", ".ppt"}:
        if suffix == ".ppt":
            raise ValueError(
                "Legacy .ppt is not supported. Open it and re-save as .pptx or export a PDF."
            )
        return _from_pptx(path)
    raise ValueError(f"Unsupported deck format: {path.suffix} (expected .pdf or .pptx)")
